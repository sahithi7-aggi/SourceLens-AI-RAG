from urllib.parse import urlparse
import io
import requests
from bs4 import BeautifulSoup
from langchain_core.documents import Document

import pytesseract

pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)


USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/131.0.0.0 Safari/537.36"
)
TIMEOUT = 20


def normalize_urls(urls):
    normalized = []
    for url in urls:
        url = url.strip()
        if not url or not url.startswith(("http://", "https://")):
            continue
        parsed = urlparse(url)
        if parsed.netloc and url not in normalized:
            normalized.append(url)
    return normalized


def _fetch_page(url):
    response = requests.get(
        url,
        headers={"User-Agent": USER_AGENT, "Accept": "text/html,application/xhtml+xml,*/*;q=0.8"},
        timeout=TIMEOUT,
        allow_redirects=True,
    )
    response.raise_for_status()
    content_type = response.headers.get("Content-Type", "").lower()
    if "text/html" not in content_type and "application/xhtml+xml" not in content_type:
        raise ValueError(f"Unsupported content type: {content_type or 'unknown'}")
    return response.text, response.url


def _extract_text(html):
    soup = BeautifulSoup(html, "html.parser")
    for element in soup(["script", "style", "noscript", "svg", "nav", "footer", "header", "aside", "form"]):
        element.decompose()
    main = soup.find("main") or soup.find("article") or soup.find(attrs={"role": "main"})
    target = main if main else soup
    lines = [line.strip() for line in target.get_text(separator="\n", strip=True).splitlines() if line.strip()]
    return "\n".join(lines)


def _format_error(exc):
    if isinstance(exc, requests.exceptions.Timeout):
        return "Request timed out while fetching the webpage."
    if isinstance(exc, requests.exceptions.ConnectionError):
        return "Could not connect to the webpage."
    if isinstance(exc, requests.exceptions.HTTPError):
        status = exc.response.status_code if exc.response is not None else None
        return f"The webpage returned HTTP {status}."
    return str(exc)


def load_urls(urls):
    documents, errors, source_results = [], [], []
    for url in normalize_urls(urls):
        try:
            html, final_url = _fetch_page(url)
            text = _extract_text(html)
            if len(text.strip()) < 100:
                raise ValueError("The webpage did not contain enough readable text.")
            documents.append(Document(
                page_content=text,
                metadata={"source": url, "final_url": final_url, "title": urlparse(final_url).netloc,
                          "document_type": "url", "extraction_method": "html"},
            ))
            source_results.append({"url": url, "status": "success", "documents": 1, "reason": ""})
        except Exception as exc:
            reason = _format_error(exc)
            errors.append({"url": url, "error": reason})
            source_results.append({"url": url, "status": "failed", "documents": 0, "reason": reason})
    return documents, errors, source_results


def _ocr_page(page, language="eng"):
    """OCR one PDF page locally. Requires Tesseract installed on the machine."""
    import pytesseract
    from PIL import Image

    # Render at a useful OCR resolution.
    pix = page.get_pixmap(dpi=220, alpha=False)
    image = Image.open(io.BytesIO(pix.tobytes("png")))
    return pytesseract.image_to_string(image, lang=language).strip()



def load_pdfs(
    uploaded_files,
    enable_ocr=True,
    ocr_language="eng",
):
    """
    Extract text from PDF files.

    When OCR is enabled, pages are OCR'd when the native PDF
    text appears empty or suspicious. This is important for
    scanned PDFs that contain a bad/garbled text layer.
    """

    import io
    import re

    import pymupdf
    import pytesseract

    from PIL import Image, ImageEnhance, ImageFilter
    from langchain_core.documents import Document

    documents = []
    errors = []
    source_results = []

    def clean_text(text):
        if not text:
            return ""

        text = text.replace("\x00", " ")
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)

        return text.strip()

    def looks_like_garbled_text(text):
        """
        Detect suspicious PDF text layers.

        A scanned PDF can contain a malformed hidden text layer.
        We don't want to trust that layer just because it is non-empty.
        """

        if not text or len(text.strip()) < 30:
            return True

        words = text.split()

        if len(words) < 5:
            return True

        # Average word length.
        avg_word_length = sum(
            len(word)
            for word in words
        ) / len(words)

        # Ratio of alphabetic characters.
        alpha_chars = sum(
            char.isalpha()
            for char in text
        )

        total_chars = max(
            len(text),
            1,
        )

        alpha_ratio = (
            alpha_chars / total_chars
        )

        # Too many tiny fragments often indicates
        # a broken PDF text layer.
        tiny_words = sum(
            len(word) <= 2
            for word in words
        )

        tiny_word_ratio = (
            tiny_words / len(words)
        )

        if avg_word_length < 2.5:
            return True

        if alpha_ratio < 0.45:
            return True

        if tiny_word_ratio > 0.65:
            return True

        return False

    def preprocess_for_ocr(image):
        """
        Improve OCR quality for scanned documents.
        """

        # Convert to grayscale.
        image = image.convert("L")

        # Increase contrast.
        image = ImageEnhance.Contrast(
            image
        ).enhance(2.0)

        # Sharpen.
        image = image.filter(
            ImageFilter.SHARPEN
        )

        return image

    for uploaded_file in uploaded_files:

        filename = uploaded_file.name

        try:
            pdf_bytes = uploaded_file.getvalue()

            pdf = pymupdf.open(
                stream=pdf_bytes,
                filetype="pdf",
            )

            file_documents = []
            ocr_pages = 0
            native_pages = 0

            for page_number, page in enumerate(
                pdf,
                start=1,
            ):

                native_text = clean_text(
                    page.get_text("text")
                )

                use_ocr = False

                if enable_ocr:

                    # OCR when native extraction is:
                    # empty OR suspicious/garbled.
                    if looks_like_garbled_text(
                        native_text
                    ):
                        use_ocr = True

                if use_ocr:

                    # 300 DPI is substantially better
                    # for OCR than a low-resolution render.
                    pix = page.get_pixmap(
                        dpi=300,
                        alpha=False,
                    )

                    image = Image.open(
                        io.BytesIO(
                            pix.tobytes("png")
                        )
                    )

                    image = preprocess_for_ocr(
                        image
                    )

                    ocr_text = pytesseract.image_to_string(
                        image,
                        lang=ocr_language,
                        config="--psm 6",
                    )

                    text = clean_text(
                        ocr_text
                    )

                    extraction_method = "ocr"

                    ocr_pages += 1

                else:

                    text = native_text

                    extraction_method = (
                        "native_pdf"
                    )

                    native_pages += 1

                if not text:
                    continue

                metadata = {
                    "source": filename,
                    "document_name": filename,
                    "document_type": "pdf",
                    "page": page_number,
                    "extraction_method": extraction_method,
                }

                file_documents.append(
                    Document(
                        page_content=text,
                        metadata=metadata,
                    )
                )

            pdf.close()

            if file_documents:

                documents.extend(
                    file_documents
                )

                source_results.append(
                    {
                        "url": filename,
                        "status": "success",
                        "reason": (
                            f"{len(file_documents)} pages extracted; "
                            f"{ocr_pages} OCR, "
                            f"{native_pages} native PDF."
                        ),
                    }
                )

            else:

                errors.append(
                    f"{filename}: No readable text found."
                )

                source_results.append(
                    {
                        "url": filename,
                        "status": "failed",
                        "reason": "No readable text found.",
                    }
                )

        except Exception as exc:

            errors.append(
                f"{filename}: {exc}"
            )

            source_results.append(
                {
                    "url": filename,
                    "status": "failed",
                    "reason": str(exc),
                }
            )

    return (
        documents,
        errors,
        source_results,
    )


def load_docx(uploaded_files):
    from docx import Document as DocxDocument
    documents, errors, source_results = [], [], []
    for uploaded_file in uploaded_files:
        filename = uploaded_file.name
        try:
            doc = DocxDocument(io.BytesIO(uploaded_file.getvalue()))
            text = "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
            if not text:
                raise ValueError("No readable text found in DOCX.")
            documents.append(Document(
                page_content=text,
                metadata={"source": filename, "document_name": filename, "document_type": "docx",
                          "extraction_method": "python-docx"},
            ))
            source_results.append({"url": filename, "status": "success", "documents": 1, "reason": ""})
        except Exception as exc:
            errors.append({"url": filename, "error": str(exc)})
            source_results.append({"url": filename, "status": "failed", "documents": 0, "reason": str(exc)})
    return documents, errors, source_results


def load_txt(uploaded_files):
    documents, errors, source_results = [], [], []
    for uploaded_file in uploaded_files:
        filename = uploaded_file.name
        try:
            text = uploaded_file.getvalue().decode("utf-8", errors="replace").strip()
            if not text:
                raise ValueError("File is empty.")
            documents.append(Document(
                page_content=text,
                metadata={"source": filename, "document_name": filename, "document_type": "txt",
                          "extraction_method": "text"},
            ))
            source_results.append({"url": filename, "status": "success", "documents": 1, "reason": ""})
        except Exception as exc:
            errors.append({"url": filename, "error": str(exc)})
            source_results.append({"url": filename, "status": "failed", "documents": 0, "reason": str(exc)})
    return documents, errors, source_results
